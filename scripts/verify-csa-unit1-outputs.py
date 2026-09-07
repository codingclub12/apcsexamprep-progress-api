#!/usr/bin/env python3
"""For each Day 1 TEACHER deck: union the class members across its code slides,
run the program, and compare to the OUTPUT panel the deck shows.

    python3 scripts/verify-csa-unit1-outputs.py <dir-of-Lesson_*-folders>

This is the check that established the Unit 1 programs are CORRECT and only the
presentation was wrong. The worked examples are split across slides, and both
halves declare the same class, so no half compiles alone. Union the members and
they run, and their real output matches what the deck prints.

Two non-matches are expected and are the decks being right rather than wrong:
1.11 prints Math.random() and says "a different number from 1 to 6 on each run"
instead of pinning a value, and 1.1 splits its three output lines across three
text boxes. 1.8 has no main at all and says so on the slide.

Needs the decks present; they are not committed, because this repository is
public and they are the paid teacher bundle.
"""
import os,re,sys,glob,subprocess,tempfile,shutil
from pptx import Presentation
NOISE='Picked up JAVA_TOOL_OPTIONS'
DECL=re.compile(r'^\s*(?:public\s+|final\s+|abstract\s+)*class\s+(\w+)',re.M)
def clean(s): return '\n'.join(l for l in (s or '').splitlines() if NOISE not in l).strip()
def body(code):
    i=code.find('{')
    if i<0: return None
    d=0
    for j in range(i,len(code)):
        if code[j]=='{': d+=1
        elif code[j]=='}':
            d-=1
            if d==0: return code[i+1:j]
    return None
work=tempfile.mkdtemp(prefix='oc'); rows=[]
os.chdir(sys.argv[1] if len(sys.argv) > 1 else '.')
for p in sorted(glob.glob('*/Day1_Deck_TEACHER.pptx')):
    prs=Presentation(p); blocks=[]; shown=None
    for sl in prs.slides:
        shapes=[sh for sh in sl.shapes if sh.has_text_frame and sh.text_frame.text.strip()]
        txt=[sh.text_frame.text for sh in shapes]
        for sh in shapes:
            if any(r.font.name=='Courier New' for pp in sh.text_frame.paragraphs for r in pp.runs):
                blocks.append(sh.text_frame.text)
        if 'OUTPUT' in txt:
            k=txt.index('OUTPUT')
            if k+1<len(txt): shown=txt[k+1].strip()
    if not blocks: continue
    # union members by class name; blocks with no decl are continuations of the previous
    byclass={}; cur=None; buf=''
    for b in blocks:
        m=DECL.search(b)
        if m:
            if cur and buf is not None:
                byclass.setdefault(cur,[]).append(buf)
            cur=m.group(1); buf=b
        else:
            buf=buf+'\n'+b if buf else b
    if cur: byclass.setdefault(cur,[]).append(buf)
    srcs={}
    for c,parts in byclass.items():
        bodies=[]
        for pt in parts:
            bb=body(pt)
            if bb is None:
                bodies=None; break
            bodies.append(bb)
        if bodies is None: srcs=None; break
        srcs[c]='public class %s\n{\n%s\n}\n'%(c,'\n'.join(bodies))
    lesson=p.split('/')[0].replace('Lesson_','')
    if not srcs: rows.append((lesson,'CANNOT_ASSEMBLE','',shown)); continue
    d=os.path.join(work,lesson); os.makedirs(d,exist_ok=True); files=[]
    for c,s in srcs.items():
        open(os.path.join(d,c+'.java'),'w').write(s); files.append(c+'.java')
    r=subprocess.run(['javac','-nowarn']+files,capture_output=True,text=True,cwd=d,timeout=120)
    if r.returncode!=0:
        e=clean(r.stderr).splitlines()
        rows.append((lesson,'COMPILE_FAIL',re.sub(r'^.*?\.java:','line ',next((l for l in e if 'error:' in l),e[0] if e else '?'))[:60],shown)); continue
    mc=next((c for c,s in srcs.items() if re.search(r'static\s+void\s+main\s*\(',s)),None)
    if not mc: rows.append((lesson,'NO_MAIN','',shown)); continue
    rr=subprocess.run(['java','-cp','.',mc],capture_output=True,text=True,cwd=d,timeout=30)
    if rr.returncode!=0:
        rows.append((lesson,'RUN_FAIL',clean(rr.stderr).splitlines()[0][:60],shown)); continue
    rows.append((lesson,'RUNS',clean(rr.stdout),shown))
shutil.rmtree(work,ignore_errors=True)
print(f"{'lesson':30} {'verdict':14} {'actual output':22} {'deck shows':22} match")
match=miss=0
for l,v,a,s in rows:
    A=' | '.join(a.split('\n')) if a else ''
    S=' | '.join((s or '').split('\n'))
    ok = (A==S) if v=='RUNS' else None
    if ok is True: match+=1
    elif v=='RUNS': miss+=1
    print(f"  {l[:28]:28} {v:14} {A[:20]:20} {S[:20]:20} {'YES' if ok else ('NO' if ok is False else '-')}")
print(f"\nprograms that assemble and run: {sum(1 for r in rows if r[1]=='RUNS')} of {len(rows)}")
print(f"printed OUTPUT matches actual: {match}   mismatched: {miss}")
sys.exit(0 if sum(1 for r in rows if r[1]=='RUNS') >= 14 else 1)
