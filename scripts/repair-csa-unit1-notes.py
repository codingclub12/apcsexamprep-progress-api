#!/usr/bin/env python3
"""
Repair the mid-word truncation in the AP CSA Unit 1 guided notes.

    python3 scripts/repair-csa-unit1-notes.py <notesdir> <deckdir> [--dry-run]

WHAT IS WRONG
Every note-taking prompt in the Unit 1 student guided notes is cut at exactly 60
characters, without regard for where a word ends. A student is handed lines like

    Classes in Java libraries are grouped into packages. A packa ______
    Every class has two kinds of members. Attributes are the dat ______
    The short version for now: a primitive holds the value itsel ______

Measured across topics 1.2 and 1.7 and systematic in both. The kit cuts at a word
boundary, which is what these should have done.

HOW IT DECIDES, AND WHY IT USES THE DECKS
Trimming every prompt back to the last space would also chop prompts that happen
to end on a real word boundary, shortening them for no reason. So the decks are
used as an ORACLE, not as a source of text: the same prose appears on the section
slides in full, so looking the prompt up there says whether the next character is
a letter, which is what makes a cut mid-word.

The repair itself never inserts text. It only removes a trailing partial word from
what is already on the page, so a lookup that goes wrong can shorten a prompt but
can never put words in a student's mouth.

No em-dashes, per repo convention.
"""

import argparse
import os
import re
import shutil
import sys

from docx import Document

try:
    from pptx import Presentation
except ImportError:
    Presentation = None

RULE = re.compile(r'^(.*?)(\s*[_ \s]{6,})$')
UNDERSCORES = re.compile(r'_{6,}')


def deck_prose(deckdir, topic):
    """Every section paragraph from that topic's decks, as one blob."""
    if Presentation is None:
        return ''
    out = []
    folder = None
    for d in os.listdir(deckdir):
        if re.match(r'Lesson_%s_' % re.escape(topic), d):
            folder = os.path.join(deckdir, d)
            break
    if folder is None:
        return ''
    for fn in sorted(os.listdir(folder)):
        if not fn.endswith('.pptx') or fn.endswith('.orig.pptx'):
            continue
        for sl in Presentation(os.path.join(folder, fn)).slides:
            for sh in sl.shapes:
                if sh.has_text_frame and sh.text_frame.text.strip():
                    out.append(sh.text_frame.text)
    return '\n'.join(out)


def trim_partial_word(prompt, prose):
    """Return the prompt cut back to a word boundary, or None to leave it."""
    p = prompt.rstrip()
    if not p or ' ' not in p:
        return None
    i = prose.find(p)
    if i < 0:
        return None                       # not found: leave it alone
    nxt = prose[i + len(p):i + len(p) + 1]
    if not nxt or not nxt.isalnum():
        return None                       # the cut is already at a boundary
    cut = p.rsplit(' ', 1)[0].rstrip()
    if not cut or cut == p:
        return None
    return cut


def repair(path, prose, dry):
    doc = Document(path)
    edits = []
    for para in doc.paragraphs:
        t = para.text
        if not UNDERSCORES.search(t):
            continue
        m = RULE.match(t)
        if not m:
            continue
        head, tail = m.group(1), m.group(2)
        if len(head.strip()) < 25:
            continue                      # a fill-in-the-blank, not a prompt
        cut = trim_partial_word(head, prose)
        if cut is None:
            continue
        edits.append((head.strip()[-28:], cut.strip()[-28:]))
        if not dry:
            #  rewrite the run that holds the prompt, leaving the rule alone
            for r in para.runs:
                if r.text.strip() and r.text.strip() in head:
                    r.text = r.text.replace(head.rstrip(), cut)
                    break
            else:
                para.runs[0].text = para.runs[0].text.replace(head.rstrip(), cut)
    if edits and not dry:
        backup = path.replace('.docx', '.orig.docx')
        if not os.path.exists(backup):
            shutil.copy2(path, backup)
        doc.save(path)
    return edits


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('notesdir')
    ap.add_argument('deckdir')
    ap.add_argument('--dry-run', action='store_true')
    a = ap.parse_args()
    total = 0
    files = 0
    cache = {}
    for root, _, names in os.walk(a.notesdir):
        for fn in sorted(names):
            if not fn.endswith('.docx') or fn.endswith('.orig.docx'):
                continue
            m = re.search(r'(\d+\.\d+)', root + '/' + fn)
            topic = m.group(1) if m else None
            if topic is None:
                continue
            if topic not in cache:
                cache[topic] = deck_prose(a.deckdir, topic)
            edits = repair(os.path.join(root, fn), cache[topic], a.dry_run)
            if not edits:
                continue
            files += 1
            print('%s  (%d prompts)' % (os.path.join(root, fn).replace(a.notesdir, '').lstrip('/'), len(edits)))
            for old, new in edits[:3]:
                print('    ...%s   ->   ...%s' % (old, new))
            total += len(edits)
    print()
    print('%d prompts trimmed across %d files%s'
          % (total, files, ' (dry run, nothing written)' if a.dry_run else ''))


if __name__ == '__main__':
    main()
