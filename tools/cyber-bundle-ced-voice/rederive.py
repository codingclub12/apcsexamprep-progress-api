"""
A SECOND implementation, deliberately unlike the first.

verify.py re-reads the output with the same extractor and the same voice.py
regexes that produced the edits. If those regexes are wrong, both agree and the
run is green for the wrong reason. That is the shape of every defect this repo
found on 2026-09-01 and 02: caught only by a check of a DIFFERENT KIND.

So this file shares nothing with the tool:

  parser     python-docx walks the OPC package and its own paragraph/table model,
             instead of the hand-rolled ElementTree walk in analyze.py
  detection  whole-word TOKEN matching against a literal word list, instead of
             voice.py's regexes. No pattern is imported from the tool.
  scope      STUDENT documents only, where no judgement is needed: every word in
             them is read by a student, so the expected count is simply zero.

If this and verify.py disagree, believe neither and go look.

  python3 rederive.py <bundle-root> <out-root>
"""
import os
import re
import sys

try:
    import docx  # python-docx
except ImportError:
    print("needs: pip install python-docx")
    raise

# Words that name the course framework rather than the security idea.
BANNED_WORDS = {'CED', 'CEDs', 'EK', 'EKs', 'LO', 'LOs'}
BANNED_PHRASES = ['college board', 'essential knowledge', 'learning objective',
                  'skill category', 'cb standard', 'cb-required', 'cb-specified',
                  'ced-listed', 'ced-based', 'ced-correct', 'ced-recommended']
# A three-part objective code: 4.1.C or 2.2.A.8. Two-part lesson numbers are fine.
CODE = re.compile(r'\b\d\.\d\.[A-F](?:\.\d+)?\b')
TOKEN = re.compile(r"[A-Za-z][A-Za-z'-]*")


def all_text(path):
    """Every paragraph python-docx can see, body and tables alike."""
    d = docx.Document(path)
    out = [p.text for p in d.paragraphs]
    for t in d.tables:
        for row in t.rows:
            for cell in row.cells:
                out.extend(p.text for p in cell.paragraphs)
    return out


def scan(path):
    hits = []
    for line in all_text(path):
        low = line.lower()
        for tok in TOKEN.findall(line):
            if tok in BANNED_WORDS:
                hits.append((tok, line))
        for ph in BANNED_PHRASES:
            if ph in low:
                hits.append((ph, line))
        for m in CODE.finditer(line):
            hits.append((m.group(0), line))
    return hits


def student_files(root):
    for dirpath, _, names in os.walk(root):
        for n in sorted(names):
            if n.endswith('.docx') and ('_STUDENT.docx' in n or n == 'Discussion.docx'):
                yield os.path.join(dirpath, n)


def main(base, out):
    before = after = 0
    files = 0
    remaining = []
    for p in student_files(out):
        files += 1
        h = scan(p)
        after += len(h)
        if h:
            remaining.append((os.path.relpath(p, out), h[:3]))
    for p in student_files(base):
        before += len(scan(p))
    print(f"student-facing documents scanned : {files}")
    print(f"framework references BEFORE      : {before}")
    print(f"framework references AFTER       : {after}")
    for rel, h in remaining[:25]:
        print(f"  STILL PRESENT {rel}")
        for tok, line in h:
            print(f"      {tok!r} in: {line.strip()[:110]}")
    return 1 if after else 0


if __name__ == '__main__':
    sys.exit(main(sys.argv[1], sys.argv[2]))
